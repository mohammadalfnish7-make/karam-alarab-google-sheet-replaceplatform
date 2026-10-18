"""One-time import of the load-plan workbook.

Empty sea and air rows are skipped. Exit rows that only have a number are skipped.
Dirty legacy cells are kept and flagged; they are not rejected.
"""

from __future__ import annotations

import re
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import AirCargo, Company, Customer, ExitRecord, ImportBatch, SeaCargo, Supplier
from app.validation import (
    bring_date_not_after_today,
    canonical_choice,
    canonical_exit_status,
    canonical_stage,
    excel_serial_to_date,
    looks_real,
    parse_compact_dmy,
    parse_dmy,
    split_invoices,
)
from app.constants import AIR_STATUSES


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _issue(field: str, message: str, raw=None) -> dict:
    item = {"field": field, "message": message}
    if raw is not None:
        item["raw"] = str(raw)
    return item


def _headers(ws, count: int) -> list[str]:
    row = next(ws.iter_rows(min_row=1, max_row=1, max_col=count))
    headers = []
    for cell in row:
        if cell.value is None:
            headers.append("")
        else:
            headers.append(str(cell.value).strip().lower())
    return headers


def _expect(headers: list[str], expected: list[str], sheet: str) -> None:
    if headers[: len(expected)] != expected:
        raise ValueError(f"{sheet} headers {headers[: len(expected)]} do not match {expected}")


def _display_datetime(value: datetime, number_format: str | None) -> str:
    fmt = (number_format or "").strip().lower()
    if fmt == "yyyy-m":
        return f"{value.year}-{value.month}"
    if fmt == "yyyy-mm":
        return f"{value.year}-{value.month:02d}"
    return value.date().isoformat()


def _label(value, field: str, number_format: str | None = None):
    if _blank(value):
        return None, [], False
    if isinstance(value, datetime):
        text = _display_datetime(value, number_format)
        return text, [_issue(field, "Excel date kept as display text", text)], False
    if isinstance(value, float):
        text = str(int(value)) if value.is_integer() else str(value)
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    else:
        text = str(value).strip()
    if not text:
        return None, [], False
    if not looks_real(text):
        return text, [_issue(field, "Placeholder value", text)], False
    return text, [], True


def _cartons(value, field: str):
    if _blank(value):
        return None, []
    if isinstance(value, bool):
        return None, [_issue(field, "Cartons value could not be read", value)]
    if isinstance(value, int) and value > 0:
        return value, []
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value), []
    text = str(value).strip()
    if re.fullmatch(r"\d+", text) and int(text) > 0:
        return int(text), []
    if re.fullmatch(r"\d+\.0+", text):
        number = int(float(text))
        if number > 0:
            return number, []
    return None, [_issue(field, "Cartons value could not be read", value)]


def _gross(value):
    if _blank(value):
        return None, False, []
    if isinstance(value, bool):
        return None, False, [_issue("gross_weight_kg", "Unparsed weight", value)]
    if isinstance(value, (int, float)):
        number = float(value)
        issues = []
        if number < 0:
            issues.append(_issue("gross_weight_kg", "Negative weight", value))
        return number, False, issues
    text = str(value).strip()
    token = text.upper()
    if token in {"NO PALLET", "NO PALLETT"}:
        issues = []
        if token == "NO PALLETT":
            issues.append(_issue("gross_weight_kg", "Normalized NO PALLETT to NO PALLET", text))
        return None, True, issues
    numeric = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", numeric):
        return float(numeric), False, []
    return None, False, [_issue("gross_weight_kg", "Unparsed weight", text)]


def _net(value):
    if _blank(value):
        return None, []
    if isinstance(value, bool):
        return None, [_issue("net_weight_kg", "Unparsed weight", value)]
    if isinstance(value, (int, float)):
        number = float(value)
        issues = []
        if number < 0:
            issues.append(_issue("net_weight_kg", "Negative weight", value))
        return number, issues
    text = str(value).strip()
    numeric = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", numeric):
        return float(numeric), []
    return None, [_issue("net_weight_kg", "Unparsed weight", text)]


def _system(value):
    if _blank(value):
        return None, []
    if isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    elif isinstance(value, int) and not isinstance(value, bool):
        raw = str(value)
    else:
        raw = str(value).strip()
    if not raw:
        return None, []
    if re.fullmatch(r"DL[0-9]+", raw, flags=re.IGNORECASE):
        return "DL" + raw[2:], []
    return raw, [_issue("system_code", "Value does not match DL followed by digits", raw)]


def _shipment(value, kind: str):
    if _blank(value):
        return None, []
    if isinstance(value, float) and value.is_integer():
        raw = str(int(value))
    elif isinstance(value, int) and not isinstance(value, bool):
        raw = str(value)
    else:
        raw = str(value).strip()
    if not raw:
        return None, []
    upper = raw.upper()
    pattern = r"D[0-9]+" if kind == "sea" else r"AR[0-9]+"
    if re.fullmatch(pattern, upper):
        return upper, []
    return upper, [_issue("shipment_number", "Value does not match the shipment pattern", raw)]


def _route(value):
    if _blank(value):
        return None, []
    raw = str(value).strip()
    token = re.sub(r"\s+", " ", raw).upper()
    if token == "TURKY":
        return "TURKEY", [_issue("shipping_route", "Normalized TURKY to TURKEY", raw)]
    if token in {"IRAN", "TURKEY"}:
        return token, []
    return token, [_issue("shipping_route", "Route is not IRAN or TURKEY", raw)]


def _stage(value):
    if _blank(value):
        return None, []
    raw = str(value).strip()
    canon = canonical_stage(raw)
    if canon is None:
        return raw, [_issue("stage", "Stage is not an allowed value", raw)]
    if raw != canon:
        return canon, [_issue("stage", f"Normalized {raw} to {canon}", raw)]
    return canon, []


def _status(value, kind: str):
    if _blank(value):
        return None, []
    raw = str(value).strip()
    canon = canonical_exit_status(raw) if kind == "exit" else canonical_choice(raw, AIR_STATUSES)
    if canon is None:
        return raw, [_issue("status", "Status is not an allowed value", raw)]
    if raw != canon:
        return canon, [_issue("status", f"Normalized {raw} to {canon}", raw)]
    return canon, []


def _as_sheet_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    parsed = parse_dmy(text)
    if parsed:
        return parsed
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _received(value):
    if _blank(value):
        return None, []
    if isinstance(value, (datetime, date)) or isinstance(value, str):
        parsed = _as_sheet_date(value)
        if parsed:
            return bring_date_not_after_today(parsed), []
    return None, [_issue("date_received", "Invalid date", value)]


def _kept_date(parsed: date | None, raw):
    if parsed is None:
        return None, [_issue("date_of_submission", "Invalid date", raw)]
    return bring_date_not_after_today(parsed), []


def _submission(value):
    if _blank(value):
        return None, []
    if isinstance(value, datetime):
        return _kept_date(value.date(), value)
    if isinstance(value, date):
        return _kept_date(value, value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return _kept_date(excel_serial_to_date(int(value)), value)
    text = str(value).strip()
    if re.fullmatch(r"\d+", text):
        return _kept_date(excel_serial_to_date(int(text)), text)
    parsed = parse_dmy(text) or parse_compact_dmy(text)
    if parsed:
        return _kept_date(parsed, text)
    try:
        return _kept_date(date.fromisoformat(text[:10]), text)
    except ValueError:
        return None, [_issue("date_of_submission", "Invalid date", text)]


def _number(value) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+", text):
        return int(text)
    return None


def _remember(bucket: dict[str, str], text: str | None, suggest: bool) -> None:
    if suggest and text:
        bucket.setdefault(text.lower(), text)


def import_workbook(db: Session, path: str | Path) -> dict:
    path = Path(path)
    filename = path.name
    existing = db.scalar(select(ImportBatch).where(ImportBatch.filename == filename))
    if existing:
        return {
            "skipped": True,
            "filename": filename,
            "sea": existing.sea_count,
            "air": existing.air_count,
            "exit": existing.exit_count,
        }

    imported_at = datetime.now(timezone.utc)
    customers: dict[str, str] = {}
    companies: dict[str, str] = {}
    suppliers: dict[str, str] = {}
    sea_rows: list[SeaCargo] = []
    air_rows: list[AirCargo] = []
    exit_rows: list[ExitRecord] = []

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        sea = workbook["SEA GOODS"]
        _expect(
            _headers(sea, 10),
            [
                "date received",
                "customer",
                "company name",
                "number of cartons",
                "gross weight",
                "net weight",
                "system code",
                "shipment number",
                "shipping route",
                "stage",
            ],
            "SEA GOODS",
        )
        for row in sea.iter_rows(min_row=2, max_col=10):
            cells = list(row)
            if all(_blank(cells[index].value) for index in range(10)):
                continue
            issues: list[dict] = []
            received, received_issues = _received(cells[0].value)
            customer, customer_issues, customer_ok = _label(cells[1].value, "customer_code", cells[1].number_format)
            company, company_issues, company_ok = _label(cells[2].value, "company_name", cells[2].number_format)
            cartons, carton_issues = _cartons(cells[3].value, "cartons")
            gross, no_pallet, gross_issues = _gross(cells[4].value)
            net, net_issues = _net(cells[5].value)
            system, system_issues = _system(cells[6].value)
            shipment, shipment_issues = _shipment(cells[7].value, "sea")
            route, route_issues = _route(cells[8].value)
            stage, stage_issues = _stage(cells[9].value)
            for group in (
                received_issues,
                customer_issues,
                company_issues,
                carton_issues,
                gross_issues,
                net_issues,
                system_issues,
                shipment_issues,
                route_issues,
                stage_issues,
            ):
                issues.extend(group)
            _remember(customers, customer, customer_ok)
            _remember(companies, company, company_ok)
            sea_rows.append(
                SeaCargo(
                    legacy=True,
                    needs_review=bool(issues),
                    issues=issues,
                    date_received=received,
                    customer_code=customer,
                    company_name=company,
                    cartons=cartons,
                    gross_weight_kg=gross,
                    gross_is_no_pallet=no_pallet,
                    net_weight_kg=net,
                    system_code=system,
                    shipment_number=shipment,
                    shipping_route=route,
                    stage=stage,
                    created_at=imported_at,
                )
            )

        air = workbook["AIR SHIPMENT"]
        _expect(
            _headers(air, 9),
            [
                "date received",
                "customer",
                "company name",
                "number of cartons",
                "gross weight",
                "net weight",
                "system code",
                "shipment number",
                "status",
            ],
            "AIR SHIPMENT",
        )
        for row in air.iter_rows(min_row=2, max_col=9):
            cells = list(row)
            if all(_blank(cells[index].value) for index in range(9)):
                continue
            issues = []
            received, received_issues = _received(cells[0].value)
            customer, customer_issues, customer_ok = _label(cells[1].value, "customer_code", cells[1].number_format)
            company, company_issues, company_ok = _label(cells[2].value, "company_name", cells[2].number_format)
            cartons, carton_issues = _cartons(cells[3].value, "cartons")
            gross, no_pallet, gross_issues = _gross(cells[4].value)
            net, net_issues = _net(cells[5].value)
            system, system_issues = _system(cells[6].value)
            shipment, shipment_issues = _shipment(cells[7].value, "air")
            status, status_issues = _status(cells[8].value, "air")
            for group in (
                received_issues,
                customer_issues,
                company_issues,
                carton_issues,
                gross_issues,
                net_issues,
                system_issues,
                shipment_issues,
                status_issues,
            ):
                issues.extend(group)
            _remember(customers, customer, customer_ok)
            _remember(companies, company, company_ok)
            air_rows.append(
                AirCargo(
                    legacy=True,
                    needs_review=bool(issues),
                    issues=issues,
                    date_received=received,
                    customer_code=customer,
                    company_name=company,
                    cartons=cartons,
                    gross_weight_kg=gross,
                    gross_is_no_pallet=no_pallet,
                    net_weight_kg=net,
                    system_code=system,
                    shipment_number=shipment,
                    status=status,
                    created_at=imported_at,
                )
            )

        exit_sheet = workbook["EXIT TRACKER"]
        _expect(
            _headers(exit_sheet, 7),
            ["no", "mark", "supplier", "invoice no", "cartons", "date of submission", "status"],
            "EXIT TRACKER",
        )
        for row in exit_sheet.iter_rows(min_row=2, max_col=7):
            cells = list(row)
            if all(_blank(cells[index].value) for index in range(1, 7)):
                continue
            number = _number(cells[0].value)
            if number is None:
                continue
            issues = []
            mark, mark_issues, mark_ok = _label(cells[1].value, "mark", cells[1].number_format)
            supplier, supplier_issues, supplier_ok = _label(cells[2].value, "supplier", cells[2].number_format)
            invoice_text, invoice_issues, _invoice_ok = _label(cells[3].value, "invoice_no", cells[3].number_format)
            invoice_numbers = split_invoices(invoice_text)
            invoice = ", ".join(invoice_numbers) if invoice_numbers else None
            cartons, carton_issues = _cartons(cells[4].value, "cartons")
            submitted, date_issues = _submission(cells[5].value)
            status, status_issues = _status(cells[6].value, "exit")
            for group in (mark_issues, supplier_issues, invoice_issues, carton_issues, date_issues, status_issues):
                issues.extend(group)
            _remember(customers, mark, mark_ok)
            _remember(suppliers, supplier, supplier_ok)
            exit_rows.append(
                ExitRecord(
                    legacy=True,
                    needs_review=bool(issues),
                    issues=issues,
                    no=number,
                    mark=mark,
                    supplier=supplier,
                    invoice_no=invoice,
                    invoice_numbers=invoice_numbers,
                    cartons=cartons,
                    date_of_submission=submitted,
                    status=status,
                    created_at=imported_at,
                )
            )
    finally:
        workbook.close()

    for code in customers.values():
        if db.scalar(select(Customer.id).where(func.lower(Customer.code) == code.lower())) is None:
            db.add(Customer(code=code, created_at=imported_at))
    for name in companies.values():
        if db.scalar(select(Company.id).where(func.lower(Company.name) == name.lower())) is None:
            db.add(Company(name=name, created_at=imported_at))
    for name in suppliers.values():
        if db.scalar(select(Supplier.id).where(func.lower(Supplier.name) == name.lower())) is None:
            db.add(Supplier(name=name, created_at=imported_at))

    db.add_all(sea_rows)
    db.add_all(air_rows)
    db.add_all(exit_rows)
    db.add(
        ImportBatch(
            filename=filename,
            sea_count=len(sea_rows),
            air_count=len(air_rows),
            exit_count=len(exit_rows),
            created_at=imported_at,
        )
    )
    db.commit()
    return {
        "skipped": False,
        "filename": filename,
        "sea": len(sea_rows),
        "air": len(air_rows),
        "exit": len(exit_rows),
    }


def _date_from_issue(item) -> date | None:
    if not isinstance(item, dict) or item.get("field") != "date_of_submission":
        return None
    raw = item.get("raw")
    if raw is None:
        return None
    text = str(raw).strip()
    parsed = parse_dmy(text) or parse_compact_dmy(text)
    if parsed:
        return parsed
    if re.fullmatch(r"\d+", text):
        return excel_serial_to_date(int(text))
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def repair_dates_after_today(db: Session) -> dict:
    """Bring sheet dates that landed after today back to the day and month that were typed."""
    today = date.today()
    corrected = 0
    for model in (SeaCargo, AirCargo):
        rows = db.scalars(select(model).where(model.date_received.is_not(None), model.date_received > today)).all()
        for row in rows:
            fixed = bring_date_not_after_today(row.date_received, today)
            if fixed != row.date_received and fixed <= today:
                row.date_received = fixed
                corrected += 1
    exit_corrected = 0
    exits = db.scalars(select(ExitRecord).where(ExitRecord.deleted_at.is_(None))).all()
    for row in exits:
        issues = list(row.issues or [])
        current = row.date_of_submission
        fixed = bring_date_not_after_today(current, today) if current and current > today else current
        kept = []
        filled = fixed
        for item in issues:
            recovered = _date_from_issue(item)
            if recovered and (filled is None or (isinstance(item, dict) and item.get("message") == "Invalid date")):
                recovered = bring_date_not_after_today(recovered, today)
                if recovered <= today:
                    filled = recovered
                    continue
            if isinstance(item, dict) and item.get("message") == "Date is after today" and filled and filled <= today:
                continue
            kept.append(item)
        if filled == current and kept == issues:
            continue
        row.date_of_submission = filled
        row.issues = kept
        row.needs_review = bool(kept)
        exit_corrected += 1
    if corrected or exit_corrected:
        db.commit()
    return {"received_corrected": corrected, "exit_corrected": exit_corrected}


def repair_invoice_numbers(db: Session) -> dict:
    """Split exit cells that list more than one invoice into invoice_numbers."""
    rows = db.scalars(select(ExitRecord)).all()
    changed = 0
    for row in rows:
        numbers = split_invoices(row.invoice_no)
        current = list(row.invoice_numbers or [])
        if row.invoice_numbers is not None and numbers == current:
            continue
        joined = ", ".join(numbers) if numbers else None
        row.invoice_numbers = numbers
        if joined != row.invoice_no:
            row.invoice_no = joined
        changed += 1
    if changed:
        db.commit()
    return {"invoices_split": changed}
