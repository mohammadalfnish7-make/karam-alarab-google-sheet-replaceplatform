"""Parse standalone container loading plan Excel (LOADING LIST + COMMODITY)."""

from __future__ import annotations

import io
import re
from typing import Any

import openpyxl

from app.errors import AppError

_CONTAINER_RE = re.compile(r"CONTAINER\s*NUMBER\s*:\s*(.+)", re.IGNORECASE)


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _intish(value: Any) -> int | None:
    value = _cell(value)
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else int(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else int(number)


def _floatish(value: Any) -> float | int | None:
    value = _cell(value)
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _text(value: Any) -> str | None:
    value = _cell(value)
    if value is None:
        return None
    return str(value).strip() or None


def parse_loading_plan(content: bytes, *, filename: str | None = None) -> dict[str, Any]:
    # Not read_only: need row_dimensions.hidden to flag rows hidden in Excel UI
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise AppError(422, {"message": f"Could not read Excel: {exc}"}) from exc

    sheet_names = {name.strip().upper(): name for name in wb.sheetnames}
    list_name = sheet_names.get("LOADING LIST")
    if not list_name:
        raise AppError(422, {"message": "Workbook must include a LOADING LIST sheet"})

    lines: list[dict[str, Any]] = []
    hidden_lines: list[dict[str, Any]] = []
    ws = wb[list_name]
    for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if index == 1:
            continue
        cells = list(row) if row else []
        while len(cells) < 10:
            cells.append(None)
        no_raw = _text(cells[0])
        name = _text(cells[4])
        cartons = _intish(cells[5])
        if no_raw and no_raw.lower().startswith("total"):
            break
        # Named marks only — avoid weight-only / blank rows inflating a mix
        if not name:
            continue
        if no_raw and not no_raw.replace(".", "", 1).isdigit() and not name:
            continue
        dim = ws.row_dimensions.get(index)
        hidden = bool(dim and dim.hidden)
        line = {
            "no": _intish(cells[0]),
            "customer_code": _text(cells[1]),
            "cust_name": _text(cells[2]),
            "mix": _intish(cells[3]),
            "name": name,
            "cartons": cartons,
            "weight": _floatish(cells[6]),
            "pw": _floatish(cells[7]),
            "total_weight": _floatish(cells[9]),
            "hidden": hidden,
            "excel_row": index,
        }
        lines.append(line)
        if hidden:
            hidden_lines.append(line)

    container_number: str | None = None
    commodities: list[dict[str, Any]] = []
    commodity_name = sheet_names.get("COMMODITY")
    if commodity_name:
        cws = wb[commodity_name]
        for index, row in enumerate(cws.iter_rows(values_only=True), start=1):
            cells = list(row) if row else []
            if index == 1:
                first = _text(cells[0] if cells else None)
                if first:
                    match = _CONTAINER_RE.search(first)
                    if match:
                        container_number = match.group(1).strip() or None
                continue
            if index == 2:
                continue
            category = _text(cells[2] if len(cells) > 2 else None)
            qty = _intish(cells[4] if len(cells) > 4 else None)
            weight = _floatish(cells[5] if len(cells) > 5 else None)
            if not category and qty is None and weight is None:
                continue
            commodities.append(
                {
                    "no": _intish(cells[0] if cells else None),
                    "hs_code": _text(cells[1] if len(cells) > 1 else None),
                    "category": category,
                    "coo": _text(cells[3] if len(cells) > 3 else None),
                    "qty": qty,
                    "weight": weight,
                }
            )

    if not lines:
        raise AppError(422, {"message": "LOADING LIST has no usable rows"})

    warnings: list[dict[str, Any]] = []
    if hidden_lines:
        sample = ", ".join(
            f"{item.get('name') or '?'} (mix {item.get('mix')})" for item in hidden_lines[:6]
        )
        more = f" +{len(hidden_lines) - 6} more" if len(hidden_lines) > 6 else ""
        warnings.append(
            {
                "code": "HIDDEN_EXCEL_ROWS",
                "message": (
                    f"{len(hidden_lines)} plan row(s) are hidden in Excel but still loaded: "
                    f"{sample}{more}"
                ),
            }
        )

    return {
        "filename": filename,
        "container_number": container_number,
        "lines": lines,
        "commodities": commodities,
        "line_count": len(lines),
        "hidden_row_count": len(hidden_lines),
        "warnings": warnings,
    }


def normalize_plan_payload(
    body: dict[str, Any] | None,
    *,
    filename: str | None = None,
) -> dict[str, Any]:
    """Normalize an in-app plan editor payload into plan_json (no Excel)."""
    if not isinstance(body, dict):
        raise AppError(422, {"message": "Plan body is required"})

    raw_lines = body.get("lines")
    if not isinstance(raw_lines, list) or not raw_lines:
        raise AppError(422, {"message": "Add at least one plan row", "field_errors": {"lines": "Required"}})

    lines: list[dict[str, Any]] = []
    field_errors: dict[str, str] = {}
    for index, raw in enumerate(raw_lines):
        if not isinstance(raw, dict):
            field_errors[f"lines.{index}"] = "Invalid row"
            continue
        name = _text(raw.get("name"))
        mix = _intish(raw.get("mix"))
        if not name:
            field_errors[f"lines.{index}.name"] = "Name is required"
            continue
        if mix is None or mix < 1:
            field_errors[f"lines.{index}.mix"] = "Mix must be 1 or greater"
            continue
        weight = _floatish(raw.get("weight"))
        total_weight = _floatish(raw.get("total_weight"))
        if total_weight is None:
            total_weight = weight
        lines.append(
            {
                "no": _intish(raw.get("no")) if raw.get("no") is not None else index + 1,
                "customer_code": _text(raw.get("customer_code")),
                "cust_name": _text(raw.get("cust_name")),
                "mix": mix,
                "name": name,
                "cartons": _intish(raw.get("cartons")),
                "weight": weight,
                "pw": _floatish(raw.get("pw")),
                "total_weight": total_weight,
            }
        )

    if field_errors:
        raise AppError(422, {"message": "Fix plan row errors", "field_errors": field_errors})
    if not lines:
        raise AppError(422, {"message": "Add at least one plan row with name and mix"})

    container_number = _text(body.get("container_number"))
    plan_name = _text(filename) or _text(body.get("filename")) or "manual"

    return {
        "filename": plan_name,
        "container_number": container_number,
        "lines": lines,
        "commodities": [],
        "line_count": len(lines),
        "hidden_row_count": 0,
        "warnings": [],
        "total_weight": sum(
            (line.get("weight") if line.get("weight") is not None else line.get("total_weight") or 0)
            for line in lines
        ),
    }
