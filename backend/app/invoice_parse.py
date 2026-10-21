"""Parse DR1249-style invoice / packing list Excel into invoice_json."""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

import openpyxl

from app.errors import AppError

_REF_RE = re.compile(r"INVOICE\s*&?\s*PAK(?:C|K)ING\s*LIST\s*:?\s*(.+)", re.IGNORECASE)
_DATE_RE = re.compile(r"DATE\s*:?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", re.IGNORECASE)
_CONT_RE = re.compile(r"(\d{4,})", re.IGNORECASE)


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def _text(value: Any) -> str | None:
    value = _cell(value)
    if value is None:
        return None
    return str(value).strip() or None


def _intish(value: Any) -> int | None:
    value = _cell(value)
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else int(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else int(number)


def _floatish(value: Any) -> float | int | None:
    value = _cell(value)
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if isinstance(value, float) and value.is_integer() else value
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _parse_date(raw: str | None) -> date | None:
    if not raw:
        return None
    text = raw.strip().replace(".", "/")
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def build_groups(lines: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for line in lines:
        mark = (line.get("mark") or "").strip()
        shop = (line.get("shop") or "").strip()
        key = (mark, shop)
        if current is None or (current["mark"], current["shop"]) != key:
            current = {
                "mark": mark or None,
                "shop": shop or None,
                "cartons": 0,
                "weight": 0.0,
                "line_count": 0,
                "has_weight": False,
            }
            groups.append(current)
        current["line_count"] += 1
        ctn = line.get("cartons")
        if isinstance(ctn, (int, float)):
            current["cartons"] += int(ctn)
        weight = line.get("weight")
        if isinstance(weight, (int, float)):
            current["weight"] += float(weight)
            current["has_weight"] = True
    out: list[dict[str, Any]] = []
    for group in groups:
        weight = group["weight"] if group["has_weight"] else None
        if weight is not None and weight == int(weight):
            weight = int(weight)
        out.append(
            {
                "mark": group["mark"],
                "shop": group["shop"],
                "cartons": group["cartons"] or None,
                "weight": weight,
                "line_count": group["line_count"],
            }
        )
    return out


def finalize_invoice_json(payload: dict[str, Any]) -> dict[str, Any]:
    lines = payload.get("lines") if isinstance(payload.get("lines"), list) else []
    groups = build_groups(lines)
    total_cartons = 0
    total_weight = 0.0
    has_weight = False
    for line in lines:
        ctn = line.get("cartons")
        if isinstance(ctn, (int, float)):
            total_cartons += int(ctn)
        weight = line.get("weight")
        if isinstance(weight, (int, float)):
            total_weight += float(weight)
            has_weight = True
    tw: float | int | None = total_weight if has_weight else None
    if tw is not None and tw == int(tw):
        tw = int(tw)
    return {
        **payload,
        "lines": lines,
        "groups": groups,
        "line_count": len(lines),
        "group_count": len(groups),
        "total_cartons": total_cartons or None,
        "total_weight": tw,
    }


def parse_invoice(content: bytes, *, filename: str | None = None) -> dict[str, Any]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        raise AppError(422, {"message": f"Could not read Excel: {exc}"}) from exc

    ws = wb[wb.sheetnames[0]]
    title = _text(ws.cell(2, 1).value) or _text(ws.cell(1, 1).value)
    plate_line = _text(ws.cell(3, 1).value)
    coo_line = _text(ws.cell(4, 1).value)

    ref = None
    if title:
        match = _REF_RE.search(title)
        if match:
            ref = match.group(1).strip() or None
        elif ":" in title:
            ref = title.split(":")[-1].strip() or None
    if not ref:
        ref = (filename or "invoice").rsplit(".", 1)[0][:64]

    invoice_date = None
    container_no = None
    if plate_line:
        date_match = _DATE_RE.search(plate_line)
        if date_match:
            invoice_date = _parse_date(date_match.group(1))
        # Prefer trailing number chunks after slash (e.g. 169 / 25279)
        parts = [p.strip() for p in plate_line.replace("|", "/").split("/") if p.strip()]
        for part in reversed(parts):
            nums = _CONT_RE.findall(part)
            if nums:
                container_no = nums[-1]
                break

    header_row = None
    for index, row in enumerate(ws.iter_rows(values_only=True, max_row=20), start=1):
        cells = [(_text(c) or "").upper() if c is not None else "" for c in (row or [])]
        if any(c == "S.NO" or c == "SNO" or c.startswith("S.NO") for c in cells[:3]) and any(
            "MARK" in c for c in cells[:5]
        ):
            header_row = index
            break
    if header_row is None:
        header_row = 5

    lines: list[dict[str, Any]] = []
    current_mark: str | None = None
    current_shop: str | None = None
    for index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if index <= header_row:
            continue
        cells = list(row) if row else []
        while len(cells) < 10:
            cells.append(None)
        no_raw = _text(cells[0])
        if no_raw and no_raw.lower().startswith("total"):
            break
        description = _text(cells[3])
        mark = _text(cells[1])
        shop = _text(cells[2])
        cartons = _intish(cells[8])
        weight = _floatish(cells[9])
        pcs = _intish(cells[5])
        if not description and cartons is None and weight is None and pcs is None and not mark:
            continue
        if mark:
            current_mark = mark
        if shop:
            current_shop = shop
        if not description and not mark and not shop:
            continue
        lines.append(
            {
                "no": _intish(cells[0]),
                "mark": current_mark,
                "shop": current_shop,
                "description": description,
                "coo": _text(cells[4]),
                "pcs": pcs,
                "hs_code": _text(cells[6]),
                "type": _text(cells[7]),
                "cartons": cartons,
                "weight": weight,
            }
        )

    if not lines:
        raise AppError(422, {"message": "Invoice sheet has no usable rows"})

    return finalize_invoice_json(
        {
            "ref": ref,
            "title": title,
            "plate_line": plate_line,
            "coo_line": coo_line,
            "container_no": container_no,
            "invoice_date": invoice_date.isoformat() if invoice_date else None,
            "filename": filename,
            "lines": lines,
        }
    )
