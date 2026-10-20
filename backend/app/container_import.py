"""Import Cont Tracker / Air Shipment Tracker / COO / Flight Schedule from the workbook.

Skips API Credentials and Form Responses 1. Idempotent by ImportBatch filename.
"""

from __future__ import annotations

import re
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    AirShipmentTrackerRow,
    ContainerRow,
    CooRow,
    FlightScheduleRow,
    ImportBatch,
)
from app.constants import CONTAINER_STAGES
from app.validation import parse_compact_dmy, parse_dmy

SKIP_SHEETS = {"API Credentials", "Form Responses 1"}


def _text(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.startswith("="):
        return None
    return text or None


def _container_stage(value) -> str | None:
    text = _text(value)
    if not text:
        return None
    normalized = re.sub(r"\s+", " ", text.strip().lower())
    for stage in CONTAINER_STAGES:
        if stage.lower() == normalized:
            return stage
    # "ЭТАП 6" / "этап6" → этап 6
    match = re.fullmatch(r"этап\s*([1-6])", normalized)
    if match:
        return f"этап {match.group(1)}"
    return text


def _int(value) -> int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    return None


def _number(value):
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value).strip().replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    parsed = parse_dmy(text) or parse_compact_dmy(text)
    if parsed:
        return parsed
    # day/month without year → assume current/recent year from workbook context
    match = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})", text)
    if match:
        day, month = int(match.group(1)), int(match.group(2))
        for year in (2026, 2025, date.today().year):
            try:
                return date(year, month, day)
            except ValueError:
                continue
    return None


def _row_values(row) -> list:
    return [cell.value for cell in row]


def _header_map(values: list) -> dict[str, int]:
    mapping = {}
    for index, value in enumerate(values):
        if value is None:
            continue
        key = str(value).strip().lower()
        if key:
            mapping[key] = index
    return mapping


def _cell(values: list, headers: dict[str, int], *names: str):
    for name in names:
        index = headers.get(name.lower())
        if index is not None and index < len(values):
            return values[index]
    return None


def _header_value(values: list, headers: dict[str, int], *, contains: str):
    needle = contains.lower()
    for key, index in headers.items():
        if needle in key and index < len(values):
            return values[index]
    return None


def _extras(values: list, used_indexes: set[int]) -> dict | None:
    extra = {}
    for index, value in enumerate(values):
        if index in used_indexes or value is None or value == "":
            continue
        text = _text(value)
        if text is None:
            continue
        extra[str(index)] = text
    return extra or None


def _has_data(values: list, start: int = 0) -> bool:
    for value in values[start:]:
        if value is None or value == "":
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return True
    return False


def import_container_tracker(db: Session, path: Path) -> dict:
    filename = f"container-tracker:{path.name}"
    existing = db.scalar(select(ImportBatch).where(ImportBatch.filename == filename))
    if existing:
        return {
            "skipped": True,
            "filename": path.name,
            "containers": existing.sea_count,
            "air_tracker": existing.air_count,
            "coo": existing.exit_count,
            "flight_schedule": 0,
        }

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    now = datetime.now(timezone.utc)
    counts = {"containers": 0, "air_tracker": 0, "coo": 0, "flight_schedule": 0}
    try:
        if "Cont Tracker" in wb.sheetnames:
            counts["containers"] = _import_containers(db, wb["Cont Tracker"], now)
        if "Air Shipment Tracker" in wb.sheetnames:
            counts["air_tracker"] = _import_air(db, wb["Air Shipment Tracker"], now)
        if "COO" in wb.sheetnames:
            counts["coo"] = _import_coo(db, wb["COO"], now)
        if "Flight Schedule" in wb.sheetnames:
            counts["flight_schedule"] = _import_flight(db, wb["Flight Schedule"], now)
    finally:
        wb.close()

    db.add(
        ImportBatch(
            filename=filename,
            sea_count=counts["containers"],
            air_count=counts["air_tracker"],
            exit_count=counts["coo"] + counts["flight_schedule"],
            created_at=now,
        )
    )
    db.commit()
    return {"skipped": False, "filename": path.name, **counts}


def _import_containers(db: Session, ws, now: datetime) -> int:
    rows = list(ws.iter_rows())
    if not rows:
        return 0
    headers = _header_map(_row_values(rows[0]))
    used = set(headers.values())
    count = 0
    for row in rows[1:]:
        values = _row_values(row)
        if not _has_data(values, start=1):
            continue
        issues = []
        raw_date = _cell(values, headers, "date")
        parsed_date = _date(raw_date)
        if raw_date not in (None, "") and parsed_date is None:
            issues.append({"field": "date", "message": "Could not parse date", "raw": _text(raw_date)})
        item = ContainerRow(
            legacy=True,
            needs_review=bool(issues),
            issues=issues,
            date=parsed_date,
            inv_no=_text(_cell(values, headers, "inv no")),
            qr=_text(_cell(values, headers, "q r")),
            previous_cont=_text(_cell(values, headers, "previous cont")),
            container_no=_text(_header_value(values, headers, contains="container")),
            kgs=_number(_cell(values, headers, "kgs")),
            carton=_int(_cell(values, headers, "carton")),
            driver_name=_text(_header_value(values, headers, contains="driver's name")),
            driver_contact=_text(_header_value(values, headers, contains="driver's contact")),
            info=_text(_header_value(values, headers, contains="info")),
            notes=_text(_cell(values, headers, "notes")),
            foreign_invoice=_text(_cell(values, headers, "foreign invoice")),
            russian_instrct=_text(_header_value(values, headers, contains="russian")),
            supplier_buyer=_text(_header_value(values, headers, contains="supplier")),
            status=_text(_cell(values, headers, "status")),
            stage=_container_stage(values[16] if len(values) > 16 else None),
            extra_json=_extras(values, used | {0, 16}),
            created_at=now,
        )
        db.add(item)
        count += 1
    db.flush()
    return count


def _import_air(db: Session, ws, now: datetime) -> int:
    rows = list(ws.iter_rows())
    if not rows:
        return 0
    headers = _header_map(_row_values(rows[0]))
    used = set(headers.values())
    count = 0
    for row in rows[1:]:
        values = _row_values(row)
        if not _has_data(values, start=1):
            continue
        issues = []
        raw_date = _cell(values, headers, "date")
        parsed_date = _date(raw_date)
        if raw_date not in (None, "") and parsed_date is None:
            issues.append({"field": "date", "message": "Could not parse date", "raw": _text(raw_date)})
        # Status is often a free date string — keep as text
        status_val = _cell(values, headers, "status")
        status_text = _text(status_val)
        if isinstance(status_val, (datetime, date)):
            status_text = status_val.date().isoformat() if isinstance(status_val, datetime) else status_val.isoformat()
        db.add(
            AirShipmentTrackerRow(
                legacy=True,
                needs_review=bool(issues),
                issues=issues,
                ref_no=_text(_cell(values, headers, "ref no")),
                vendor=_text(_cell(values, headers, "vendor")),
                airline=_text(_cell(values, headers, "airline")),
                date=parsed_date,
                item=_text(_cell(values, headers, "item")),
                cartons=_int(_cell(values, headers, "crtn")),
                weight_kg=_number(_cell(values, headers, "wt(kg)")),
                awb=_text(_cell(values, headers, "awb")),
                status=status_text,
                comment=_text(_cell(values, headers, "comment")),
                extra_json=_extras(values, used | {0}),
                created_at=now,
            )
        )
        count += 1
    db.flush()
    return count


def _import_coo(db: Session, ws, now: datetime) -> int:
    rows = list(ws.iter_rows())
    if not rows:
        return 0
    headers = _header_map(_row_values(rows[0]))
    used = set(headers.values())
    count = 0
    for row in rows[1:]:
        values = _row_values(row)
        if not _has_data(values):
            continue
        db.add(
            CooRow(
                legacy=True,
                needs_review=False,
                issues=[],
                sr=_int(_cell(values, headers, "sr")),
                inv_number=_text(_cell(values, headers, "inv number")),
                container_no=_text(_cell(values, headers, "container no")),
                kjs=_text(_cell(values, headers, "kjs")),
                grg=_text(_cell(values, headers, "grg")),
                k_exit=_text(_cell(values, headers, "k-exit")),
                extra_json=_extras(values, used),
                created_at=now,
            )
        )
        count += 1
    db.flush()
    return count


def _import_flight(db: Session, ws, now: datetime) -> int:
    rows = list(ws.iter_rows())
    if not rows:
        return 0
    headers = _header_map(_row_values(rows[0]))
    # Only use left-block columns 0-7 for primary fields
    primary_keys = {"sr", "airline", "date", "day", "narrow", "wide", "status", "column 1"}
    used = {index for key, index in headers.items() if key in primary_keys and index < 8}
    count = 0
    for row in rows[1:]:
        values = _row_values(row)
        left = values[:8]
        if not _has_data(left):
            # side-block only rows still stored if they have right-side data
            if not _has_data(values, start=8):
                continue
        issues = []
        raw_date = values[2] if len(values) > 2 else None
        parsed_date = _date(raw_date)
        if raw_date not in (None, "") and parsed_date is None and _has_data(left):
            issues.append({"field": "date", "message": "Could not parse date", "raw": _text(raw_date)})
        db.add(
            FlightScheduleRow(
                legacy=True,
                needs_review=bool(issues),
                issues=issues,
                sr=_int(values[0] if values else None),
                airline=_text(values[1] if len(values) > 1 else None),
                date=parsed_date,
                day=_text(values[3] if len(values) > 3 else None),
                narrow=_text(values[4] if len(values) > 4 else None),
                wide=_text(values[5] if len(values) > 5 else None),
                status=_text(values[6] if len(values) > 6 else None),
                column_1=_text(values[7] if len(values) > 7 else None),
                extra_json=_extras(values, used),
                created_at=now,
            )
        )
        count += 1
    db.flush()
    return count
